
import csv
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"

# Find the AIHW workbook by its worksheet name.
care_file = None

for path in DATA.glob("*.xlsx"):
    workbook = load_workbook(path, read_only=True, data_only=True)

    if "Table 4.1 (LGA)" in workbook.sheetnames:
        care_file = path
        workbook.close()
        break

    workbook.close()

if care_file is None:
    raise FileNotFoundError(
        "Could not find an Excel workbook containing "
        "'Table 4.1 (LGA)' in the data folder."
    )

print("=" * 60)
print("AIHW RESIDENTIAL-CARE DATA")
print("=" * 60)
print("Workbook:", care_file.name)

workbook = load_workbook(
    care_file, read_only=True, data_only=True
)
sheet = workbook["Table 4.1 (LGA)"]

print("\nFirst 8 worksheet rows:")
for row in sheet.iter_rows(
    min_row=1, max_row=8, values_only=True
):
    print(row[:5])

# Find the ABS LGA dataset.
abs_file = DATA / "lga_ageing_2025.csv"

if not abs_file.exists():
    raise FileNotFoundError(f"Missing file: {abs_file}")

print("\n" + "=" * 60)
print("ABS LGA AGEING DATA")
print("=" * 60)

with abs_file.open(
    "r", newline="", encoding="utf-8-sig"
) as file:
    reader = csv.DictReader(file)

    print("Columns:", reader.fieldnames)

    print("\nFirst 5 records:")
    for i, row in enumerate(reader):
        if i == 5:
            break

        print({
            key: row[key]
            for key in [
                "state", "lga_code", "lga_name",
                "total_population", "age_65_plus",
                "pct_65_plus"
            ]
            if key in row
        })

workbook.close()

print("\nInspection complete.")
