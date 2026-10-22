import csv
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
HISTORY_FILE = ROOT / "data" / "australia_age_history_1971_2025.csv"
PROJECTION_FILE = ROOT / "data" / "abs_population_projections_australia_low_2022_2071.xlsx"
OUTPUT_FILE = ROOT / "data" / "australia_ageing_observed_projected_1971_2071.csv"


def percentage(count, total):
    return round(100 * count / total, 6) if total else 0.0


def read_history():
    records = []
    with HISTORY_FILE.open("r", newline="", encoding="utf-8-sig") as file:
        reader = csv.DictReader(file)
        required = {
            "year", "total_population", "age_65_plus", "age_85_plus",
            "pct_65_plus", "pct_85_plus"
        }
        missing = required - set(reader.fieldnames or [])
        if missing:
            raise ValueError(f"Historical CSV is missing columns: {sorted(missing)}")

        for row in reader:
            records.append({
                "year": int(row["year"]),
                "total_population": int(float(row["total_population"])),
                "age_65_plus": int(float(row["age_65_plus"])),
                "age_85_plus": int(float(row["age_85_plus"])),
                "pct_65_plus": float(row["pct_65_plus"]),
                "pct_85_plus": float(row["pct_85_plus"]),
                "period": "Observed",
            })

    if not records:
        raise ValueError(f"No historical records found in {HISTORY_FILE}")

    records.sort(key=lambda item: item["year"])
    years = [item["year"] for item in records]
    if len(years) != len(set(years)):
        raise ValueError("Historical CSV contains duplicate years.")
    if years != list(range(years[0], years[-1] + 1)):
        raise ValueError("Historical CSV has missing years.")
    return records


def extract_abs_persons_series():
    if not PROJECTION_FILE.exists():
        raise FileNotFoundError(f"Projection workbook not found: {PROJECTION_FILE}")

    workbook = load_workbook(PROJECTION_FILE, read_only=True, data_only=True)
    population_by_year = defaultdict(dict)
    all_ages = set()

    try:
        for sheet_name in ("Data1", "Data2"):
            worksheet = workbook[sheet_name]
            header = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
            person_columns = []

            for column_index, label in enumerate(header):
                if not isinstance(label, str):
                    continue
                match = re.search(
                    r"Projected persons\s*;\s*Persons\s*;\s*(.*?)\s*;\s*Series",
                    label,
                    re.IGNORECASE,
                )
                if not match:
                    continue

                age_label = match.group(1).strip()
                if age_label.lower().startswith("100"):
                    age = 100
                else:
                    try:
                        age = int(age_label)
                    except ValueError:
                        continue

                person_columns.append((column_index, age))
                all_ages.add(age)

            print(f"{sheet_name}: found {len(person_columns)} Persons age series")

            for row in worksheet.iter_rows(values_only=True):
                date_value = row[0] if row else None
                if not isinstance(date_value, (date, datetime)):
                    continue

                year = date_value.year
                for column_index, age in person_columns:
                    if column_index >= len(row):
                        raise ValueError(f"Missing column for age {age} in {sheet_name}, {year}")
                    value = row[column_index]
                    if value is None:
                        raise ValueError(f"Missing value in {sheet_name}, year {year}, age {age}")
                    if age in population_by_year[year]:
                        raise ValueError(f"Duplicate Persons series for age {age} in {year}")
                    population_by_year[year][age] = int(round(float(value)))
    finally:
        workbook.close()

    expected_ages = set(range(101))
    if all_ages != expected_ages:
        missing = sorted(expected_ages - all_ages)
        unexpected = sorted(all_ages - expected_ages)
        raise ValueError(f"Unexpected age series. Missing: {missing}; unexpected: {unexpected}")

    expected_years = set(range(2022, 2072))
    if set(population_by_year) != expected_years:
        found = sorted(population_by_year)
        raise ValueError(f"Expected years 2022–2071; found {found[0]}–{found[-1]}")

    for year, counts in population_by_year.items():
        if set(counts) != expected_ages:
            missing = sorted(expected_ages - set(counts))
            raise ValueError(f"Year {year} is missing age series: {missing}")

    print("Age-series validation: PASS (ages 0–100+)")
    print("Projection-year validation: PASS (2022–2071)")
    return population_by_year


def calculate_projection_records(population_by_year):
    records = {}
    for year, age_counts in sorted(population_by_year.items()):
        total_population = sum(age_counts.values())
        age_65_plus = sum(value for age, value in age_counts.items() if age >= 65)
        age_85_plus = sum(value for age, value in age_counts.items() if age >= 85)
        if total_population <= 0:
            raise ValueError(f"Invalid total population in {year}")

        records[year] = {
            "year": year,
            "total_population": total_population,
            "age_65_plus": age_65_plus,
            "age_85_plus": age_85_plus,
            "pct_65_plus": percentage(age_65_plus, total_population),
            "pct_85_plus": percentage(age_85_plus, total_population),
            "period": "Projected",
        }
    return records


def main():
    historical = read_history()
    last_observed_year = max(item["year"] for item in historical)
    print(f"Historical observations: {historical[0]['year']}–{last_observed_year}")

    population_by_year = extract_abs_persons_series()
    projected_records = calculate_projection_records(population_by_year)

    # Preserve historical values through the final observed year.
    # Append ABS low-series projections only for subsequent years.
    future_projections = [
        record for year, record in sorted(projected_records.items())
        if year > last_observed_year
    ]
    combined = sorted(historical + future_projections, key=lambda item: item["year"])

    actual_years = [item["year"] for item in combined]
    if actual_years != list(range(actual_years[0], actual_years[-1] + 1)):
        raise ValueError("Combined dataset has missing or duplicate years.")

    columns = [
        "year", "total_population", "age_65_plus", "age_85_plus",
        "pct_65_plus", "pct_85_plus", "period"
    ]
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    with OUTPUT_FILE.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(file, fieldnames=columns)
        writer.writeheader()
        writer.writerows(combined)

    last_observed = next(item for item in historical if item["year"] == last_observed_year)
    projected_2071 = projected_records[2071]

    print("\n" + "=" * 55)
    print("CHART 11 DATA PREPARATION COMPLETE")
    print("=" * 55)
    print(f"Created: {OUTPUT_FILE}")
    print(f"Total annual records: {len(combined)}")
    print(f"Observed: {historical[0]['year']}–{last_observed_year}")
    print(f"Projected: {last_observed_year + 1}–2071")
    print(f"\nObserved {last_observed_year}:")
    print(f"  Population: {last_observed['total_population']:,}")
    print(f"  Aged 65+: {last_observed['pct_65_plus']:.2f}%")
    print(f"  Aged 85+: {last_observed['pct_85_plus']:.2f}%")

    if last_observed_year in projected_records:
        baseline = projected_records[last_observed_year]
        print(f"\nABS low-series {last_observed_year} (comparison only):")
        print(f"  Aged 65+: {baseline['pct_65_plus']:.2f}%")
        print(f"  Aged 85+: {baseline['pct_85_plus']:.2f}%")

    print("\nABS low-series projection for 2071:")
    print(f"  Population: {projected_2071['total_population']:,}")
    print(f"  Aged 65+: {projected_2071['pct_65_plus']:.2f}%")
    print(f"  Aged 85+: {projected_2071['pct_85_plus']:.2f}%")
    print("\nAll checks passed.")


if __name__ == "__main__":
    main()
