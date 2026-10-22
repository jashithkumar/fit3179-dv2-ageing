
import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
ABS_FILE = DATA / "lga_ageing_2025.csv"
OUTPUT_FILE = DATA / "state_care_vs_85plus_2025.csv"


def normalise_code(value):
    """Convert numeric or string LGA codes to comparable strings."""
    if value is None:
        return None

    text = str(value).strip()

    if not text:
        return None

    if text.endswith(".0"):
        text = text[:-2]

    if text.isdigit():
        return str(int(text))

    return text


def to_number(value):
    """Convert numeric cells; return None for missing/suppressed values."""
    if value is None:
        return None

    if isinstance(value, (int, float)):
        return int(value) if float(value).is_integer() else float(value)

    text = str(value).strip().replace(",", "")

    if text.lower() in {"", "n.p.", "n.a.", "na", "-"}:
        return None

    try:
        number = float(text)
        return int(number) if number.is_integer() else number
    except ValueError:
        return None


# --------------------------------------------------
# 1. Read ABS population data and build LGA mapping
# --------------------------------------------------

if not ABS_FILE.exists():
    raise FileNotFoundError(f"Missing file: {ABS_FILE}")

abs_by_code = {}
population_85_by_state = defaultdict(int)
duplicate_abs_codes = []

with ABS_FILE.open(
    "r", newline="", encoding="utf-8-sig"
) as file:
    reader = csv.DictReader(file)

    required = {
        "state", "lga_code", "lga_name", "age_85_plus"
    }

    if not required.issubset(reader.fieldnames or []):
        raise ValueError(
            f"ABS CSV is missing required columns: "
            f"{required - set(reader.fieldnames or [])}"
        )

    for row in reader:
        code = normalise_code(row["lga_code"])

        if not code:
            continue

        if code in abs_by_code:
            duplicate_abs_codes.append(code)
            continue

        state = row["state"].strip()
        age_85 = to_number(row["age_85_plus"])

        if age_85 is None:
            raise ValueError(
                f"Missing age_85_plus for "
                f"{row['lga_name']} ({code})"
            )

        abs_by_code[code] = {
            "state": state,
            "lga_name": row["lga_name"],
            "age_85_plus": age_85,
        }

        population_85_by_state[state] += age_85

if duplicate_abs_codes:
    raise ValueError(
        f"Duplicate ABS LGA codes: {duplicate_abs_codes[:20]}"
    )

print(f"ABS LGAs loaded: {len(abs_by_code)}")


# --------------------------------------------------
# 2. Find and read the AIHW workbook
# --------------------------------------------------

care_file = None

for path in DATA.glob("*.xlsx"):
    workbook = load_workbook(
        path, read_only=True, data_only=True
    )

    if "Table 4.1 (LGA)" in workbook.sheetnames:
        care_file = path
        workbook.close()
        break

    workbook.close()

if care_file is None:
    raise FileNotFoundError(
        "No workbook containing 'Table 4.1 (LGA)' "
        "was found in the data folder."
    )

print(f"AIHW workbook: {care_file.name}")

workbook = load_workbook(
    care_file, read_only=True, data_only=True
)

try:
    sheet = workbook["Table 4.1 (LGA)"]
    rows = sheet.iter_rows(values_only=True)

    header_row = None
    header_index = None

    for index, row in enumerate(rows):
        values = [
            str(value).strip().lower()
            if value is not None else ""
            for value in row
        ]

        if "lga code" in values and "total" in values:
            header_row = values
            header_index = index
            break

    if header_row is None:
        raise ValueError(
            "Could not locate the LGA code and Total headers."
        )

    code_col = header_row.index("lga code")
    name_col = header_row.index("lga name")
    total_col = header_row.index("total")

    care_by_state = defaultdict(int)
    seen_codes = set()
    unmatched = []
    missing_totals = []
    duplicate_care_codes = []
    care_rows = 0

    # Continue after the header row.
    for row in sheet.iter_rows(
        min_row=header_index + 2, values_only=True
    ):
        if code_col >= len(row):
            continue

        code = normalise_code(row[code_col])

        if not code:
            continue

        name = (
            str(row[name_col]).strip()
            if name_col < len(row) and row[name_col] is not None
            else "Unknown LGA"
        )

        care_rows += 1

        if code in seen_codes:
            duplicate_care_codes.append((code, name))
            continue

        seen_codes.add(code)

        if code not in abs_by_code:
            unmatched.append((code, name))
            continue

        total = (
            to_number(row[total_col])
            if total_col < len(row)
            else None
        )

        if total is None:
            missing_totals.append((code, name))
            continue

        care_by_state[abs_by_code[code]["state"]] += total

finally:
    workbook.close()


# --------------------------------------------------
# 3. Validate matching before creating output
# --------------------------------------------------

print(f"AIHW LGA rows read: {care_rows}")
print(f"Matched LGA codes: {len(seen_codes & set(abs_by_code))}")
print(f"Unmatched AIHW LGAs: {len(unmatched)}")
print(f"Missing AIHW totals: {len(missing_totals)}")
print(f"Duplicate AIHW codes: {len(duplicate_care_codes)}")

if unmatched:
    print("\nFirst unmatched LGAs:")
    for code, name in unmatched[:15]:
        print(f"  {code}: {name}")

if missing_totals:
    print("\nFirst LGAs with missing totals:")
    for code, name in missing_totals[:15]:
        print(f"  {code}: {name}")

if duplicate_care_codes:
    print("\nDuplicate AIHW codes:")
    for code, name in duplicate_care_codes[:15]:
        print(f"  {code}: {name}")

if unmatched or missing_totals or duplicate_care_codes:
    raise SystemExit(
        "\nValidation failed. No chart dataset was written. "
        "Review the mismatches before proceeding."
    )


# --------------------------------------------------
# 4. Aggregate by state and calculate national shares
# --------------------------------------------------

total_85_plus = sum(population_85_by_state.values())
total_care_users = sum(care_by_state.values())

if total_85_plus <= 0 or total_care_users <= 0:
    raise ValueError("A national denominator is zero or invalid.")

states = sorted(
    set(population_85_by_state) | set(care_by_state)
)

results = []

for state in states:
    population_85 = population_85_by_state.get(state, 0)
    care_users = care_by_state.get(state, 0)

    population_share = population_85 / total_85_plus * 100
    care_share = care_users / total_care_users * 100

    results.append({
        "state": state,
        "age_85_plus_population": population_85,
        "age_85_plus_share_pct": round(population_share, 4),
        "residential_care_users_service_location": care_users,
        "residential_care_share_pct": round(care_share, 4),
        "difference_pp": round(care_share - population_share, 4),
    })


# --------------------------------------------------
# 5. Write chart dataset
# --------------------------------------------------

fields = [
    "state",
    "age_85_plus_population",
    "age_85_plus_share_pct",
    "residential_care_users_service_location",
    "residential_care_share_pct",
    "difference_pp",
]

with OUTPUT_FILE.open(
    "w", newline="", encoding="utf-8"
) as file:
    writer = csv.DictWriter(file, fieldnames=fields)
    writer.writeheader()
    writer.writerows(results)


# --------------------------------------------------
# 6. Print summary for manual verification
# --------------------------------------------------

print("\n" + "=" * 65)
print("CHART 12 DATA PREPARATION COMPLETE")
print("=" * 65)
print(f"85+ population denominator: {total_85_plus:,}")
print(f"Residential-care users denominator: {total_care_users:,}")
print(f"Output: {OUTPUT_FILE}")
print("\nState-level results:")

for row in results:
    print(
        f"{row['state']:<35} "
        f"85+ share: {row['age_85_plus_share_pct']:>6.2f}% | "
        f"Care share: {row['residential_care_share_pct']:>6.2f}% | "
        f"Difference: {row['difference_pp']:>+6.2f} pp"
    )

print("\nSum of 85+ shares:",
      round(sum(r["age_85_plus_share_pct"] for r in results), 2))
print("Sum of care shares:",
      round(sum(r["residential_care_share_pct"] for r in results), 2))
